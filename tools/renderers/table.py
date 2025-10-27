# table.py
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

def render(slide, data, geom: dict, context: dict):
    """
    セルごとのスタイル（fill/text_color/bold/italic/underline/align/v_align/font_size）を
    YAMLから直接指定できるテーブル描画。行ゼブラは任意（セル指定があれば上書き）。
    """

    # --- data 正規化 ---
    if hasattr(data, "model_dump"):
        data = data.model_dump()
    data = data or {}

    def _rgb(hex_or_none, fallback=None):
        s = (hex_or_none or "").strip()
        if not s:
            return fallback
        if s.startswith("#"):
            s = s[1:]
        if len(s) != 6:
            return fallback
        try:
            return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))
        except Exception:
            return fallback

    # 既定スタイル
    header_fill_rgb = _rgb(data.get("header_fill") or "#DDEBF7")
    header_text_rgb = _rgb(data.get("header_text_color") or "#0F172A")
    band_fill_rgb   = _rgb(data.get("band_fill") or "#F2F2F2")
    cell_text_rgb   = _rgb(data.get("cell_text_color") or "#111111")
    table_style_name = data.get("table_style")

    font_size        = int(data.get("font_size", 11))
    header_font_size = int(data.get("header_font_size", 12))
    align_default    = (data.get("align") or "left").lower()
    col_align_cfg    = list(data.get("col_align") or [])
    wrap             = bool(data.get("wrap", True))
    v_align_default  = (data.get("vertical_align") or "middle").lower()
    cell_padding_pt  = int(data.get("cell_padding_pt", 4))

    banding          = bool(data.get("banding", True))
    band_start_idx   = int(data.get("band_start_index", 0))

    column_headers   = list(data.get("column_headers") or [])
    row_headers      = list(data.get("row_headers") or [])
    rows             = [list(r) for r in (data.get("rows") or [])]

    col_widths       = data.get("col_widths")
    row_heights      = data.get("row_heights")

    # 方向解決
    def _pp_align(s: str):
        return {
            "left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT
        }.get((s or "").lower(), PP_ALIGN.LEFT)

    def _pp_v_anchor(s: str):
        return {
            "top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM
        }.get((s or "").lower(), MSO_ANCHOR.MIDDLE)

    v_anchor_default = _pp_v_anchor(v_align_default)

    # セル仕様の解釈（スカラー or {value, style...}）
    def _unpack_cell(spec):
        # 返り値: (text, style_dict)
        # style_dict: {"fill": RGB, "text_color": RGB, "bold":bool, "italic":bool, "underline":bool,
        #              "align": PP_ALIGN|None, "v_anchor": MSO_ANCHOR|None, "font_size": int|None}
        if isinstance(spec, dict):
            val = spec.get("value", "")
            # style ショートハンドも許容
            st = dict(spec.get("style") or {})
            for k in ("fill", "text_color", "bold", "italic", "underline", "align", "vertical_align", "font_size"):
                if k in spec and spec[k] is not None:
                    st[k] = spec[k]
            style = {
                "fill": _rgb(st.get("fill"), None),
                "text_color": _rgb(st.get("text_color"), None),
                "bold": st.get("bold", None),
                "italic": st.get("italic", None),
                "underline": st.get("underline", None),
                "align": _pp_align(st.get("align")) if st.get("align") else None,
                "v_anchor": _pp_v_anchor(st.get("vertical_align")) if st.get("vertical_align") else None,
                "font_size": int(st["font_size"]) if st.get("font_size") else None,
            }
            return ("" if val is None else str(val), style)
        # スカラー
        return ("" if spec is None else str(spec), {})

    # データ形状
    has_col_header = len(column_headers) > 0
    has_row_header = len(row_headers) > 0
    max_cols_in_rows = max((len(r) for r in rows), default=0)
    data_cols = max(max_cols_in_rows, len(column_headers))
    n_cols = data_cols + (1 if has_row_header else 0)
    n_rows = len(rows) + (1 if has_col_header else 0)

    if has_row_header and len(row_headers) < len(rows):
        row_headers += [""] * (len(rows) - len(row_headers))
    elif has_row_header and len(row_headers) > len(rows):
        row_headers = row_headers[:len(rows)]

    # データ列揃え
    for i in range(len(rows)):
        if len(rows[i]) < data_cols:
            rows[i].extend([""] * (data_cols - len(rows[i])))
        elif len(rows[i]) > data_cols:
            rows[i] = rows[i][:data_cols]

    shapes = context.get("shapes_target", slide.shapes)

    # 生成
    table_shape = shapes.add_table(n_rows, n_cols, geom["left"], geom["top"], geom["width"], geom["height"])
    tbl = table_shape.table
    if table_style_name:
        try:
            tbl.style = table_style_name
        except Exception:
            pass
    # 列幅・行高（重み配分）
    def _apply_col_widths():
        if not col_widths: return
        weights = [max(float(w), 0.0) for w in col_widths]
        if len(weights) != n_cols: return
        total = sum(weights) or 1.0
        for j in range(n_cols):
            w = int(geom["width"] * (weights[j] / total))
            try: tbl.columns[j].width = w
            except Exception: pass

    def _apply_row_heights():
        if not row_heights: return
        weights = [max(float(h), 0.0) for h in row_heights]
        if len(weights) != n_rows: return
        total = sum(weights) or 1.0
        for i in range(n_rows):
            h = int(geom["height"] * (weights[i] / total))
            try: tbl.rows[i].height = h
            except Exception: pass

    _apply_col_widths()
    _apply_row_heights()

    # 列寄せ 既定
    default_align = _pp_align(align_default)
    col_align_resolved = [ _pp_align(col_align_cfg[j]) if j < len(col_align_cfg) else default_align for j in range(n_cols) ]

    def _apply_text(cell, text, *, is_header=False, col_index=0,
                    text_color=None, bold=None, italic=None, underline=None,
                    align=None, v_anchor=None, font_size_override=None):
        tf = cell.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = "" if text is None else str(text)
        # 文字色/サイズ
        p.font.color.rgb = (text_color if text_color is not None else (header_text_rgb if is_header else cell_text_rgb))
        p.font.size = Pt(header_font_size if is_header and font_size_override is None else
                         (font_size_override if font_size_override is not None else (header_font_size if is_header else font_size)))
        # 装飾
        if bold is None:
            p.font.bold = True if is_header else False
        else:
            p.font.bold = bool(bold)
        if italic is not None:
            p.font.italic = bool(italic)
        if underline is not None:
            p.font.underline = bool(underline)
        # 寄せ
        p.alignment = align if align is not None else (PP_ALIGN.CENTER if is_header else col_align_resolved[col_index])
        # 縦寄せ
        cell.vertical_anchor = v_anchor if v_anchor is not None else v_anchor_default
        # 折返し・パディング
        tf.word_wrap = bool(wrap)
        try:
            pad = Pt(cell_padding_pt)
            tf.margin_left = pad; tf.margin_right = pad; tf.margin_top = pad; tf.margin_bottom = pad
        except Exception:
            pass

    def _fill(cell, rgb):
        if not rgb: return
        try:
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb
        except Exception:
            pass

    # セルスタイルの保持（banding後に上書き適用するため）
    per_cell_styles = {}

    # 上ヘッダー
    if has_col_header:
        for j in range(n_cols):
            cell = tbl.cell(0, j)
            if has_row_header and j == 0:
                _apply_text(cell, "", is_header=True, col_index=0)
                _fill(cell, header_fill_rgb)
                continue
            data_col_idx = j - (1 if has_row_header else 0)
            text, style = _unpack_cell(column_headers[data_col_idx] if data_col_idx < len(column_headers) else "")
            _apply_text(cell, text, is_header=True, col_index=j,
                        text_color=style.get("text_color"), bold=style.get("bold"),
                        italic=style.get("italic"), underline=style.get("underline"),
                        align=style.get("align"), v_anchor=style.get("v_anchor"),
                        font_size_override=style.get("font_size"))
            # 既定ヘッダー背景
            _fill(cell, header_fill_rgb)
            # ヘッダーセル個別fillがあれば上書き
            if style.get("fill"):
                _fill(cell, style["fill"])

    # 左ヘッダー
    if has_row_header:
        for i in range(n_rows):
            if has_col_header and i == 0:  # 左上
                continue
            data_row_idx = i - (1 if has_col_header else 0)
            text, style = _unpack_cell(row_headers[data_row_idx] if data_row_idx < len(row_headers) else "")
            cell = tbl.cell(i, 0)
            _apply_text(cell, text, is_header=True, col_index=0,
                        text_color=style.get("text_color"), bold=style.get("bold"),
                        italic=style.get("italic"), underline=style.get("underline"),
                        align=style.get("align"), v_anchor=style.get("v_anchor"),
                        font_size_override=style.get("font_size"))
            _fill(cell, header_fill_rgb)
            if style.get("fill"):
                _fill(cell, style["fill"])

    # データ本体（テキストだけ先に入れて、スタイルは保持）
    for i in range(len(rows)):
        for j in range(data_cols):
            ti = i + (1 if has_col_header else 0)
            tj = j + (1 if has_row_header else 0)
            text, style = _unpack_cell(rows[i][j])
            cell = tbl.cell(ti, tj)
            _apply_text(cell, text, is_header=False, col_index=tj,
                        text_color=style.get("text_color"), bold=style.get("bold"),
                        italic=style.get("italic"), underline=style.get("underline"),
                        align=style.get("align"), v_anchor=style.get("v_anchor"),
                        font_size_override=style.get("font_size"))
            per_cell_styles[(ti, tj)] = style  # fillは banding 後に適用

    # 行ゼブラ（データ領域のみ）
    if banding:
        start_row = 1 if has_col_header else 0
        start_col = 1 if has_row_header else 0
        for i in range(start_row, n_rows):
            band_idx = (i - start_row)
            if (band_idx - band_start_idx) % 2 == 0:
                for j in range(start_col, n_cols):
                    _fill(tbl.cell(i, j), band_fill_rgb)

    # セル個別の fill を最後に適用（banding を上書き）
    for (ti, tj), st in per_cell_styles.items():
        if st.get("fill"):
            _fill(tbl.cell(ti, tj), st["fill"])

    return table_shape
