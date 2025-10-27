from __future__ import annotations

from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt
from pptx.dml.color import RGBColor

from ..utils import normalize_color, create_chip_with_text
from ..schemas.comparison_panel import ComparisonPanelSchema

EMU_PER_INCH = 914400
PT_PER_INCH = 72.0


def _px_from_pt(pt: float) -> int:
    # 1pt ~ 1/72inch → EMU
    return int(Pt(pt))


def _text_box(
    shapes,
    left,
    top,
    width,
    height,
    text: str,
    pt: int,
    bold=False,
    color="#000000",
    align=PP_ALIGN.LEFT,
):
    tb = shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    try:
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = 0
    p.space_after = 0
    p.text = text.strip()
    f = p.runs[0].font if p.runs else p.font
    f.size = Pt(pt)
    f.bold = bold
    r, g, b = normalize_color(color)
    f.color.rgb = RGBColor(r, g, b)
    return tb


def _chip(
    shapes,
    left,
    top,
    width,
    height,
    text: str,
    bg_hex: str,
    radius_pct: float,
    pt: int,
    text_color=None,
):
    """枠線なし角丸ボックス + テキスト"""
    return create_chip_with_text(
        shapes, left, top, width, height, text, bg_hex, radius_pct, pt, text_color
    )


def render(slide, data: ComparisonPanelSchema, geom: dict, context: dict):
    """
    比較パネル（左右または上下比較）。
    direction = "horizontal" の場合は左右配置、"vertical" の場合は上下配置。
    geom は EMU 完成形: {"left","top","width","height"}。
    """
    # Convert the validated Pydantic model back to dict for processing
    data_dict = data.model_dump()

    shapes = context.get("shapes_target", slide.shapes)
    theme = context.get("theme", {}) or {}

    # 入力と既定
    left = geom["left"]
    top = geom["top"]
    width = geom["width"]
    height = geom["height"]
    title = (data_dict or {}).get("title", "") or ""
    direction = (data_dict or {}).get("direction", "horizontal")
    style = (data_dict or {}).get("style", {}) or {}

    # Style settings (with backward compatibility for old field names)
    section_gap_pct = float(
        style.get("section_gap_pct", style.get("column_gap_pct", 3))
    )
    padding_pct = float(style.get("padding_pct", 4))
    panel_title_pt = int(style.get("panel_title_pt", 18))
    section_title_pt = int(style.get("section_title_pt", style.get("col_title_pt", 16)))
    item_pt = int(style.get("item_pt", 14))
    chip_radius_pct = float(style.get("chip_radius_pct", 0.30))
    chip_vpad_pt = float(style.get("chip_vpad_pt", 6))
    chip_spacing_pt = float(style.get("chip_spacing_pt", 6))

    # 算出（EMU）
    pad = int(width * (padding_pct / 100.0))
    gap = int(width * (section_gap_pct / 100.0))

    inner_left = left + pad
    inner_top = top + pad
    inner_width = width - pad * 2
    inner_height = height - pad * 2

    # 全体タイトル領域
    title_h = 0
    if title:
        title_h = int(Pt(panel_title_pt * 1.6))  # 行高ゆとり分を含める
        _text_box(
            shapes,
            inner_left,
            inner_top,
            inner_width,
            title_h,
            title,
            panel_title_pt,
            bold=True,
            color="#000000",
            align=PP_ALIGN.LEFT,
        )

    # セクション配置の計算
    content_top = inner_top + title_h
    content_height = inner_height - title_h

    if direction == "vertical":
        # 縦並び（上下配置）
        _render_vertical_layout(
            shapes,
            data_dict,
            theme,
            inner_left,
            content_top,
            inner_width,
            content_height,
            gap,
            section_title_pt,
            item_pt,
            chip_radius_pct,
            chip_vpad_pt,
            chip_spacing_pt,
        )
    else:
        # 横並び（左右配置） - 既存の処理
        _render_horizontal_layout(
            shapes,
            data_dict,
            theme,
            inner_left,
            content_top,
            inner_width,
            content_height,
            gap,
            section_title_pt,
            item_pt,
            chip_radius_pct,
            chip_vpad_pt,
            chip_spacing_pt,
        )


def _render_horizontal_layout(
    shapes,
    data_dict,
    theme,
    inner_left,
    content_top,
    inner_width,
    content_height,
    gap,
    section_title_pt,
    item_pt,
    chip_radius_pct,
    chip_vpad_pt,
    chip_spacing_pt,
):
    """横並び（左右配置）のレンダリング - 既存の処理"""
    left_spec = (data_dict or {}).get("left", {}) or {}
    right_spec = (data_dict or {}).get("right", {}) or {}

    col_width = (inner_width - gap) // 2

    # セクションヘッダ高さ
    has_titles = left_spec.get("title") or right_spec.get("title")
    section_head_h = int(Pt(section_title_pt * 1.5)) if has_titles else 0

    # チップ行の寸法
    chip_h = int(Pt(item_pt + chip_vpad_pt * 2))

    # セクション領域の原点
    left_x = inner_left
    right_x = inner_left + col_width + gap

    # チップの色（黒いテキストが読みやすい薄い色をデフォルトに）
    left_chip_color = left_spec.get("chip_color") or "#F0F8FF"  # さらに薄い青色

    right_chip_color = right_spec.get("chip_color") or "#F8FCFF"  # 非常に薄い青色

    # 左ヘッダ
    y_cursor_left = content_top
    if left_spec.get("title"):
        _text_box(
            shapes,
            left_x,
            y_cursor_left,
            col_width,
            section_head_h,
            left_spec["title"],
            section_title_pt,
            bold=True,
            color="#000000",
            align=PP_ALIGN.LEFT,
        )
        y_cursor_left += section_head_h + Pt(4)

    # 右ヘッダ
    y_cursor_right = content_top
    if right_spec.get("title"):
        _text_box(
            shapes,
            right_x,
            y_cursor_right,
            col_width,
            section_head_h,
            right_spec["title"],
            section_title_pt,
            bold=True,
            color="#000000",
            align=PP_ALIGN.LEFT,
        )
        y_cursor_right += section_head_h + Pt(4)

    # 左チップ群
    for text in left_spec.get("items", []):
        _chip(
            shapes,
            left_x,
            int(y_cursor_left),
            col_width,
            chip_h,
            text=text,
            bg_hex=left_chip_color,
            radius_pct=chip_radius_pct,
            pt=item_pt,
        )
        y_cursor_left += chip_h + Pt(chip_spacing_pt)

    # 右チップ群
    for text in right_spec.get("items", []):
        _chip(
            shapes,
            right_x,
            int(y_cursor_right),
            col_width,
            chip_h,
            text=text,
            bg_hex=right_chip_color,
            radius_pct=chip_radius_pct,
            pt=item_pt,
        )
        y_cursor_right += chip_h + Pt(chip_spacing_pt)


def _render_vertical_layout(
    shapes,
    data_dict,
    theme,
    inner_left,
    content_top,
    inner_width,
    content_height,
    gap,
    section_title_pt,
    item_pt,
    chip_radius_pct,
    chip_vpad_pt,
    chip_spacing_pt,
):
    """縦並び（上下配置）のレンダリング"""
    top_spec = (data_dict or {}).get("top", {}) or {}
    bottom_spec = (data_dict or {}).get("bottom", {}) or {}

    section_height = (content_height - gap) // 2

    # セクションヘッダ高さ
    has_titles = top_spec.get("title") or bottom_spec.get("title")
    section_head_h = int(Pt(section_title_pt * 1.5)) if has_titles else 0

    # チップ行の寸法
    chip_h = int(Pt(item_pt + chip_vpad_pt * 2))

    # セクション領域の原点
    top_y = content_top
    bottom_y = content_top + section_height + gap

    # チップの色（黒いテキストが読みやすい薄い色をデフォルトに）
    top_chip_color = top_spec.get("chip_color") or "#F0F8FF"  # さらに薄い青色

    bottom_chip_color = bottom_spec.get("chip_color") or "#F8FCFF"  # 非常に薄い青色

    # 上部ヘッダ
    y_cursor_top = top_y
    if top_spec.get("title"):
        _text_box(
            shapes,
            inner_left,
            y_cursor_top,
            inner_width,
            section_head_h,
            top_spec["title"],
            section_title_pt,
            bold=True,
            color="#000000",
            align=PP_ALIGN.LEFT,
        )
        y_cursor_top += section_head_h + Pt(4)

    # 下部ヘッダ
    y_cursor_bottom = bottom_y
    if bottom_spec.get("title"):
        _text_box(
            shapes,
            inner_left,
            y_cursor_bottom,
            inner_width,
            section_head_h,
            bottom_spec["title"],
            section_title_pt,
            bold=True,
            color="#000000",
            align=PP_ALIGN.LEFT,
        )
        y_cursor_bottom += section_head_h + Pt(4)

    # 上部チップ群
    for text in top_spec.get("items", []):
        _chip(
            shapes,
            inner_left,
            int(y_cursor_top),
            inner_width,
            chip_h,
            text=text,
            bg_hex=top_chip_color,
            radius_pct=chip_radius_pct,
            pt=item_pt,
        )
        y_cursor_top += chip_h + Pt(chip_spacing_pt)

    # 下部チップ群
    for text in bottom_spec.get("items", []):
        _chip(
            shapes,
            inner_left,
            int(y_cursor_bottom),
            inner_width,
            chip_h,
            text=text,
            bg_hex=bottom_chip_color,
            radius_pct=chip_radius_pct,
            pt=item_pt,
        )
        y_cursor_bottom += chip_h + Pt(chip_spacing_pt)
